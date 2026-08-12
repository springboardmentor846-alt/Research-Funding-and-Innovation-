import io
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate_pdf_report(profile, score_data, publications, patents, funding_matches, commercialization):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle", parent=styles["Title"], fontSize=20, spaceAfter=6
    )
    heading_style = ParagraphStyle(
        "HeadingStyle", parent=styles["Heading2"], spaceBefore=16, spaceAfter=8
    )
    body_style = styles["BodyText"]

    elements = []

    elements.append(Paragraph("Innovation Intelligence Report", title_style))
    elements.append(
        Paragraph(
            f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            body_style,
        )
    )
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Profile Summary", heading_style))
    elements.append(Paragraph(f"Research Domains: {profile.research_domains or 'N/A'}", body_style))
    elements.append(Paragraph(f"Organization: {profile.organization_name or 'N/A'}", body_style))

    elements.append(Paragraph("Innovation Score", heading_style))
    elements.append(
        Paragraph(
            f"<b>{score_data['innovation_score']} / 100</b> — {score_data['rating']}",
            body_style,
        )
    )

    breakdown_rows = [["Factor", "Score", "Weight"]]
    for key, val in score_data["breakdown"].items():
        breakdown_rows.append([key.replace("_", " ").title(), str(val["score"]), val["weight"]])

    breakdown_table = Table(breakdown_rows, colWidths=[7 * cm, 4 * cm, 4 * cm])
    breakdown_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#131B34")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")]),
            ]
        )
    )
    elements.append(breakdown_table)

    elements.append(Paragraph(f"Publications ({len(publications)})", heading_style))
    if publications:
        for p in publications[:10]:
            elements.append(Paragraph(f"• {p.title} ({p.year or 'N/A'})", body_style))
    else:
        elements.append(Paragraph("No publications on record.", body_style))

    elements.append(Paragraph(f"Patents ({len(patents)})", heading_style))
    if patents:
        for pt in patents[:10]:
            elements.append(Paragraph(f"• {pt.title} — {pt.assignee or 'N/A'}", body_style))
    else:
        elements.append(Paragraph("No patents on record.", body_style))

    elements.append(Paragraph(f"Matched Funding Opportunities ({len(funding_matches)})", heading_style))
    if funding_matches:
        for f in funding_matches[:10]:
            elements.append(Paragraph(f"• {f.title} — {f.amount}", body_style))
    else:
        elements.append(Paragraph("No matched funding opportunities.", body_style))

    elements.append(Paragraph("Commercialization Recommendations", heading_style))
    if commercialization:
        for rec in commercialization:
            elements.append(
                Paragraph(f"• [{rec['priority'].upper()}] {rec['type']}: {rec['recommendation']}", body_style)
            )
    else:
        elements.append(Paragraph("No recommendations available yet.", body_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_excel_report(profile, score_data, publications, patents, funding_matches):
    wb = Workbook()

    header_fill = PatternFill(start_color="131B34", end_color="131B34", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Field", "Value"])
    for cell in ws_summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws_summary.append(["Research Domains", profile.research_domains or "N/A"])
    ws_summary.append(["Organization", profile.organization_name or "N/A"])
    ws_summary.append(["Innovation Score", score_data["innovation_score"]])
    ws_summary.append(["Rating", score_data["rating"]])
    ws_summary.column_dimensions["A"].width = 22
    ws_summary.column_dimensions["B"].width = 40

    ws_pub = wb.create_sheet("Publications")
    ws_pub.append(["Title", "Authors", "Year", "Source"])
    for cell in ws_pub[1]:
        cell.fill = header_fill
        cell.font = header_font
    for p in publications:
        ws_pub.append([p.title, p.authors or "", p.year or "", p.source or ""])
    for col, width in zip("ABCD", [50, 30, 10, 20]):
        ws_pub.column_dimensions[col].width = width

    ws_pat = wb.create_sheet("Patents")
    ws_pat.append(["Title", "Assignee", "Filing Date", "Patent Number"])
    for cell in ws_pat[1]:
        cell.fill = header_fill
        cell.font = header_font
    for pt in patents:
        ws_pat.append([pt.title, pt.assignee or "", pt.filing_date or "", pt.patent_number or ""])
    for col, width in zip("ABCD", [50, 30, 15, 20]):
        ws_pat.column_dimensions[col].width = width

    ws_fund = wb.create_sheet("Funding Matches")
    ws_fund.append(["Title", "Source", "Amount", "Deadline"])
    for cell in ws_fund[1]:
        cell.fill = header_fill
        cell.font = header_font
    for f in funding_matches:
        ws_fund.append([f.title, f.source or "", f.amount or "", f.deadline or ""])
    for col, width in zip("ABCD", [50, 25, 20, 15]):
        ws_fund.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer